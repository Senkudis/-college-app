# -*- coding: utf-8 -*-
import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

# الألوان
NAVY  = hex_to_rgb('0D3B66')
BLUE  = hex_to_rgb('2176C7')
AMBER = hex_to_rgb('F59E0B')
INK   = hex_to_rgb('10161F')
MUTED = hex_to_rgb('5A6B7D')
LIGHT = hex_to_rgb('F2F5F9')
WHITE = hex_to_rgb('FFFFFF')
GREEN = hex_to_rgb('1B6B47')
CARD_BORDER = hex_to_rgb('DCE3EB')
SKY_LIGHT = hex_to_rgb('CFE0F2')
RED_ACCENT = hex_to_rgb('B84A3A')
PURPLE = hex_to_rgb('6B4FA8')

FONT = 'Arial'

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

W, H, M = 13.333, 7.5, 0.7

def add_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def set_para(p, text, font_size, color, bold=False, align=PP_ALIGN.RIGHT, rtl=True):
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    if rtl:
        pPr = p._p.get_or_add_pPr()
        pPr.set('rtl', '1')

def head(slide, title, kicker=None):
    if kicker:
        tx = slide.shapes.add_textbox(Inches(M), Inches(0.42), Inches(W - M * 2), Inches(0.35))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        set_para(p, kicker, 13, BLUE, bold=True)
    
    y = 0.72 if kicker else 0.55
    tx = slide.shapes.add_textbox(Inches(M), Inches(y), Inches(W - M * 2), Inches(0.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    set_para(p, title, 32, NAVY, bold=True)

def card(slide, x, y, w, h, title, body="", icon=None, tone=WHITE, border_color=CARD_BORDER, accent=BLUE, title_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = tone
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()

    # Icon circle if needed
    if icon:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + w - 0.95), Inches(y + 0.22), Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = accent
        circle.line.fill.background()
        tf = circle.text_frame
        p = tf.paragraphs[0]
        set_para(p, icon, 18, WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)

    # Title & Body in one textbox
    tx_w = w - (1.15 if icon else 0.5)
    tx = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.22), Inches(tx_w), Inches(h - 0.44))
    tf = tx.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    set_para(p0, title, 16, title_color, bold=True)
    
    if body:
        p1 = tf.add_paragraph()
        p1.space_before = Pt(8)
        body_color = SKY_LIGHT if tone != WHITE and tone != LIGHT else MUTED
        set_para(p1, body, 12.5, body_color)

def stat_box(slide, x, y, w, num, label, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1)
    
    tx = slide.shapes.add_textbox(Inches(x), Inches(y + 0.15), Inches(w), Inches(1.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    set_para(p0, num, 50, color, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    p1 = tf.add_paragraph()
    p1.space_before = Pt(4)
    set_para(p1, label, 13, MUTED, align=PP_ALIGN.CENTER)

# ══════════════ 1. الغلاف ══════════════
s1 = prs.slides.add_slide(blank_layout)
add_bg(s1, NAVY)

tx = s1.shapes.add_textbox(Inches(M), Inches(1.3), Inches(W - M * 2), Inches(3.2))
tf = tx.text_frame
tf.word_wrap = True
p0 = tf.paragraphs[0]
set_para(p0, 'كلية النصر التقنية — جامعة النصر التقنية', 16, hex_to_rgb('9FC2E8'), bold=True)

p1 = tf.add_paragraph()
p1.space_before = Pt(12)
set_para(p1, 'النظام الجامعي الرقمي المتكامل', 46, WHITE, bold=True)

p2 = tf.add_paragraph()
p2.space_before = Pt(16)
set_para(p2, 'منصة واحدة تدير الدراسة والحضور والتقييم والمالية والتواصل — على الويب وتطبيق الهاتف', 17, SKY_LIGHT)

# أرقام الغلاف
items = [('٥', 'أدوار'), ('٣٠', 'وحدة'), ('٥٠', 'شاشة'), ('٢.١.٠', 'الإصدار')]
for i, (n, l) in enumerate(items):
    x = W - M - 1.6 - i * 1.8
    tx_stat = s1.shapes.add_textbox(Inches(x), Inches(5.0), Inches(1.6), Inches(1.0))
    tf_s = tx_stat.text_frame
    p_num = tf_s.paragraphs[0]
    set_para(p_num, n, 28, AMBER, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    p_lbl = tf_s.add_paragraph()
    set_para(p_lbl, l, 12, hex_to_rgb('9FC2E8'), align=PP_ALIGN.CENTER)

tx_foot = s1.shapes.add_textbox(Inches(M), Inches(6.6), Inches(W - M * 2), Inches(0.4))
set_para(tx_foot.text_frame.paragraphs[0], 'عرض تقديمي للإدارة الأكاديمية', 12, hex_to_rgb('7FA5CC'))

# ══════════════ 2. لماذا ══════════════
s2 = prs.slides.add_slide(blank_layout)
head(s2, 'ما الذي يحلّه هذا النظام', 'الدافع')
problems = [
    ('الورق والواتساب', 'المحاضرات والواجبات والإعلانات موزّعة بين مجموعات دردشة وأوراق — تضيع ولا تُؤرشف.'),
    ('الحضور اليدوي', 'كشوف ورقية تُحسب يدوياً، ولا إنذار مبكر للطالب المتعثّر قبل فوات الأوان.'),
    ('المالية المنفصلة', 'الرسوم في دفاتر مستقلة عن السجل الأكاديمي، والطالب لا يعرف رصيده إلا بالسؤال.'),
    ('لا مصدر واحد', 'كل جهة تحتفظ ببياناتها، فلا تقرير موحّد يعتمد عليه القرار الأكاديمي.')
]
for i, (t, b) in enumerate(problems):
    col = i % 2
    row = i // 2
    card(s2, M + col * 6.15, 1.85 + row * 2.25, 5.8, 2.0, t, b, icon='✕', accent=RED_ACCENT)

# ══════════════ 3. نظرة عامة بالأرقام ══════════════
s3 = prs.slides.add_slide(blank_layout)
add_bg(s3, LIGHT)
head(s3, 'النظام في أرقامه', 'نظرة عامة')
stats = [('٣٠', 'وحدة خدمية (API)', BLUE), ('٥٠', 'شاشة تفاعلية', NAVY), ('٢٨', 'نموذج بيانات', GREEN), ('٥', 'أدوار بصلاحيات', AMBER)]
for i, (n, l, c) in enumerate(stats):
    stat_box(s3, M + i * 3.05, 1.95, 2.85, n, l, color=c)

card(s3, M, 4.25, 12.0, 2.3, 'بنية واحدة تخدم الويب والهاتف معاً',
     'الخادم Node.js مع قاعدة MongoDB يخدم موقع الويب وتطبيق الأندرويد وتطبيق الآيفون من نفس المصدر — أي تحديث يصل الجميع فوراً دون إعادة بناء. التطبيق مبني بـ Capacitor مع إشعارات فورية وبصمة ودعم العمل بلا إنترنت.',
     icon='◆', accent=NAVY)

# ══════════════ 4. الأدوار ══════════════
s4 = prs.slides.add_slide(blank_layout)
head(s4, 'خمسة أدوار — كل دور يرى ما يخصّه فقط', 'الصلاحيات')
roles = [
    ('الطالب', 'المواد، المحاضرات، الواجبات والتسليم، الحضور، الاختبارات، الرسوم، السجل الأكاديمي.', BLUE),
    ('قائد الدفعة', 'كل صلاحيات الطالب، بالإضافة إلى رفع المحاضرات ورصد الحضور وإرسال تنبيهات دفعته.', PURPLE),
    ('أستاذ المادة', 'إدارة مواده ومحاضراته، إنشاء الاختبارات، تقييم التسليمات، رصد الحضور، تقارير دفعاته.', GREEN),
    ('العميد', 'صلاحية كاملة: المستخدمون، الدفعات، المالية، التقارير، البلاغات، إعدادات النظام.', NAVY),
    ('مسؤول النظام', 'صلاحيات العميد التقنية مع مراقبة تشغيلية ومؤشرات أداء الخادم.', MUTED)
]
for i, (t, b, c) in enumerate(roles):
    y = 1.75 + i * 1.05
    circ = s4.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W - M - 0.65), Inches(y + 0.12), Inches(0.55), Inches(0.55))
    circ.fill.solid()
    circ.fill.fore_color.rgb = c
    circ.line.fill.background()
    set_para(circ.text_frame.paragraphs[0], str(i + 1), 16, WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    
    tx = s4.shapes.add_textbox(Inches(M), Inches(y), Inches(11.2), Inches(0.9))
    tf = tx.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    set_para(p0, t, 16, NAVY, bold=True)
    p1 = tf.add_paragraph()
    set_para(p1, b, 12.5, MUTED)

# ══════════════ 5. المحور الأكاديمي ══════════════
s5 = prs.slides.add_slide(blank_layout)
add_bg(s5, LIGHT)
head(s5, 'الدورة الأكاديمية الكاملة', 'المحور الأول')
flow = [
    ('١', 'المواد والدفعات', 'العميد ينشئ الدفعات والمواد ويسند الأساتذة.'),
    ('٢', 'المحاضرات', 'رفع الشيتات والفيديو مع ترتيب معتمد وحفظ للعمل بلا إنترنت.'),
    ('٣', 'الواجبات', 'تكليف بموعد نهائي، تسليم الطالب بملف، ومنع التسليم بعد الموعد.'),
    ('٤', 'التقييم', 'الأستاذ يرصد الدرجة والملاحظات، والطالب يُشعَر فوراً.')
]
for i, (n, t, b) in enumerate(flow):
    card(s5, M + i * 3.05, 1.85, 2.85, 2.5, t, b, icon=n, accent=BLUE)

card(s5, M, 4.65, 5.8, 1.9, 'إعادة التسليم بمهلة محددة',
     'الطالب الذي حصل على أقل من ٦٠٪ يمكنه إعادة التسليم خلال سبعة أيام من الموعد، وتُوسَم النسخة كمتأخرة أمام المقيّم.',
     icon='↺', accent=AMBER)
card(s5, M + 6.15, 4.65, 5.85, 1.9, 'متابعة من لم يسلّم',
     'الأستاذ يرى قائمة الطلاب الذين لم يسلّموا التكليف مفلترة حسب الدفعة، فيتابعهم قبل إغلاق الموعد.',
     icon='⚑', accent=GREEN)

# ══════════════ 6. الحضور والإنذار المبكر ══════════════
s6 = prs.slides.add_slide(blank_layout)
head(s6, 'الحضور والإنذار المبكر', 'المحور الثاني')
card(s6, M, 1.8, 5.8, 2.15, 'رصد الحضور الإلكتروني',
     'الأستاذ أو قائد الدفعة يرصد الحضور لكل مادة في كل يوم. سجل واحد لكل حصة يمنع الازدواج، وتُحفظ هوية من رصده.',
     icon='✓', accent=GREEN)
card(s6, M + 6.15, 1.8, 5.85, 2.15, 'تنبيه تلقائي عند الخطر',
     'حين تنزل نسبة حضور الطالب تحت ٧٥٪ يصله تنبيه فوري بنسبته الحالية والحد المطلوب — قبل أن يفقد حقه في الامتحان.',
     icon='!', accent=RED_ACCENT)
card(s6, M, 4.2, 12.0, 2.35, 'تقرير الإنذار المبكر — للعميد والأستاذ',
     'يجمع التقرير ثلاثة مؤشرات لكل طالب: تراجع الحضور، وتراكم الواجبات غير المسلَّمة، وتدنّي نتائج الاختبارات — ويصنّف الطلاب إلى خطر مرتفع ومتوسط ومنخفض.\nالغرض: تدخّل الكلية مبكراً بدل اكتشاف التعثّر في نتيجة نهاية الفصل.',
     icon='◈', accent=NAVY)

# ══════════════ 7. التقييم الإلكتروني ══════════════
s7 = prs.slides.add_slide(blank_layout)
add_bg(s7, LIGHT)
head(s7, 'الاختبارات وبنك الامتحانات', 'المحور الثالث')
feats = [
    ('اختبارات موقوتة', 'الأستاذ يحدد نافذة الإتاحة ومدة الاختبار. المؤقّت يبدأ عند فتح الطالب للاختبار ويُغلق تلقائياً عند انتهاء المدة.', '⏱', BLUE),
    ('تصحيح فوري', 'أسئلة اختيار من متعدد بدرجات مخصصة لكل سؤال، والتصحيح والدرجة يظهران لحظة التسليم.', '✍', GREEN),
    ('بنك الامتحانات السابقة', 'أرشيف امتحانات الفاينل مصنّف بالمادة والسنة والفصل، متاح للطلاب للمراجعة.', '📄', NAVY),
    ('محاكاة الامتحان', 'الطالب يؤدي امتحاناً سابقاً محاكاةً فيحصل على نتيجته ومراجعة لكل سؤال وتشخيص لمستواه.', '◎', AMBER)
]
for i, (t, b, ic, c) in enumerate(feats):
    col = i % 2
    row = i // 2
    card(s7, M + col * 6.15, 1.85 + row * 2.35, 5.8, 2.1, t, b, icon=ic, accent=c)

# ══════════════ 8. كيدي ══════════════
s8 = prs.slides.add_slide(blank_layout)
add_bg(s8, NAVY)
tx = s8.shapes.add_textbox(Inches(M), Inches(0.55), Inches(11.9), Inches(1.8))
tf = tx.text_frame
tf.word_wrap = True
set_para(tf.paragraphs[0], 'المساعد الذكي الأكاديمي', 13, hex_to_rgb('7FB3E8'), bold=True)
p1 = tf.add_paragraph()
set_para(p1, '«كيدي» — يقرأ محتوى المادة فعلياً', 34, WHITE, bold=True)
p2 = tf.add_paragraph()
p2.space_before = Pt(8)
set_para(p2, 'ليس روبوت محادثة عاماً: كيدي يستخرج نص محاضرات المادة من ملفاتها ويجيب منها هي، فتكون الإجابة من منهج الكلية لا من الإنترنت.', 15, SKY_LIGHT)

ai_feats = [
    ('التلخيص والشرح', 'يلخّص محاضرة أو المادة كاملة، ويشرح أي جزئية بلغة مبسّطة.'),
    ('توليد بطاقات الاستذكار', 'يحوّل المحاضرة إلى بطاقات مراجعة سريعة تلقائياً.'),
    ('توليد اختبارات تجريبية', 'يبني أسئلة متوقعة من محتوى المادة لمراجعة ما قبل الامتحان.'),
    ('مصحّح أخطاء البرمجة', 'يشرح سبب خطأ الكود ويقترح التصحيح لطلاب تقنية المعلومات.')
]
for i, (t, b) in enumerate(ai_feats):
    col = i % 2
    row = i // 2
    card(s8, M + col * 6.15, 2.85 + row * 1.65, 5.8, 1.45, t, b, tone=hex_to_rgb('164B7D'), border_color=hex_to_rgb('256296'), title_color=WHITE)

tx_note = s8.shapes.add_textbox(Inches(M), Inches(6.35), Inches(11.9), Inches(0.4))
set_para(tx_note.text_frame.paragraphs[0], 'يعمل على أربعة مزوّدين بالتناوب (Gemini · Qwen · Cohere · Groq) — تعطُّل أحدهم لا يوقف الخدمة.', 12, hex_to_rgb('8FB8DF'))

# ══════════════ 9. أدوات الطالب ══════════════
s9 = prs.slides.add_slide(blank_layout)
head(s9, 'أدوات تدعم الطالب خارج المحاضرة', 'المحور الرابع')
tools = [
    ('مختبر البرمجة السحابي', 'كتابة كود C++ وتشغيله من المتصفح أو الهاتف دون تثبيت أي برنامج.', '⌨', BLUE),
    ('غرف المذاكرة الجماعية', 'غرفة لكل دفعة بمهام مشتركة ومؤقّت بومودورو يتزامن بين الطلاب لحظياً.', '◷', PURPLE),
    ('بطاقات الاستذكار', 'بطاقات مراجعة سريعة تُولَّد من المحاضرة أو يكتبها الطالب بنفسه.', '▤', GREEN),
    ('لوحة المتصدرين', 'نقاط تُمنح على إنجاز الواجبات وأداء الاختبارات، وترتيب يحفّز المنافسة.', '★', AMBER),
    ('السجل الأكاديمي الموثق', 'سجل رقمي بالحضور والدرجات والإنجازات، برقم تسلسلي وبصمة تحقق.', '◫', NAVY),
    ('تقييم جودة المقررات', 'استبيان مجهول الهوية يقيس رضا الطلاب عن كل مقرر وأستاذه.', '✉', MUTED)
]
for i, (t, b, ic, c) in enumerate(tools):
    col = i % 3
    row = i // 3
    card(s9, M + col * 4.05, 1.85 + row * 2.35, 3.8, 2.1, t, b, icon=ic, accent=c)

# ══════════════ 10. المالية ══════════════
s10 = prs.slides.add_slide(blank_layout)
add_bg(s10, LIGHT)
head(s10, 'الوحدة المالية', 'المحور الخامس')
card(s10, M, 1.85, 3.8, 2.3, 'الرسوم الجامعية',
     'العميد يضيف رسوماً لطالب واحد أو لدفعة كاملة، بخصم وتاريخ استحقاق، ويُشعَر الطالب فوراً.', icon='₤', accent=NAVY)
card(s10, M + 4.05, 1.85, 3.8, 2.3, 'السداد والتخصيص التلقائي',
     'كل دفعة تُخصَّم من أقدم رسم مستحق أولاً (FIFO)، وتُسجَّل العملية كاملة في معاملة واحدة لا تنقسم.', icon='⇄', accent=GREEN)
card(s10, M + 8.1, 1.85, 3.9, 2.3, 'إشعارات وإيصالات السداد',
     'الطالب يرفع صورة إيصال التحويل من هاتفه، والإدارة تراجعه وتعتمده أو ترفضه بسبب مكتوب.', icon='⎙', accent=BLUE)
card(s10, M, 4.45, 12.0, 2.1, 'سجل الأرصدة وتقرير المتأخرات',
     'لوحة واحدة تعرض إجمالي الرسوم والمدفوع والرصيد لكل طالب، مع تقرير بالمتأخرات وتذكير تلقائي للطلاب المتأخرين عن السداد. الأرقام مشتقة من مصدر واحد فلا تتضارب بين الشاشات.', icon='▦', accent=AMBER)

# ══════════════ 11. التواصل ══════════════
s11 = prs.slides.add_slide(blank_layout)
head(s11, 'التواصل والإشعارات', 'المحور السادس')
comm = [
    ('إشعارات فورية على الهاتف', 'تصل الهاتف فوراً — درجة جديدة، تنبيه غياب، رسوم مستحقة، إعلان عاجل.', '⚡'),
    ('توجيه دقيق', 'الإشعار يُوجَّه لدفعة أو فئة أو شخص بعينه، ولا يصل من لا يعنيه.', '◉'),
    ('التقويم الأكاديمي الموحد', 'المواعيد والامتحانات والفعاليات في تقويم واحد يراه الطالب والأستاذ.', '📅'),
    ('أسئلة على الفيديو', 'الطالب يسأل عند لحظة محددة في المحاضرة المسجّلة، فيجيبه الأستاذ في موضعها.', '💬')
]
for i, (t, b, ic) in enumerate(comm):
    y = 1.8 + i * 1.05
    circ = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W - M - 0.65), Inches(y + 0.08), Inches(0.55), Inches(0.55))
    circ.fill.solid()
    circ.fill.fore_color.rgb = BLUE
    circ.line.fill.background()
    set_para(circ.text_frame.paragraphs[0], ic, 16, WHITE, align=PP_ALIGN.CENTER, rtl=False)
    
    tx = s11.shapes.add_textbox(Inches(M), Inches(y), Inches(11.2), Inches(0.9))
    tf = tx.text_frame
    tf.word_wrap = True
    set_para(tf.paragraphs[0], t, 16, NAVY, bold=True)
    p1 = tf.add_paragraph()
    set_para(p1, b, 12.5, MUTED)

card(s11, M, 6.0, 12.0, 0.95, 'الإبلاغ عن المحتوى المسيء وحظر المستخدمين — مع التزام بالمراجعة خلال ٢٤ ساعة',
     icon='⚑', tone=LIGHT, border_color=CARD_BORDER, accent=GREEN)

# ══════════════ 12. الأمان ══════════════
s12 = prs.slides.add_slide(blank_layout)
add_bg(s12, LIGHT)
head(s12, 'الأمان وخصوصية البيانات', 'الحوكمة')
sec = [
    ('كلمات المرور', 'مشفّرة بـ bcrypt ولا تُخزَّن كنص. تغييرها إلزامي عند أول دخول.'),
    ('الدخول بالبصمة', 'دعم البصمة والتعرّف على الوجه عبر Passkeys بدل كتابة كلمة المرور.'),
    ('جلسة واحدة نشطة', 'الدخول من جهاز جديد ينهي الجلسة السابقة تلقائياً.'),
    ('عزل الدفعات', 'الخادم يمنع الوصول لبيانات دفعة أخرى — لا الواجهة فقط.'),
    ('ملفات محمية', 'تسليمات الطلاب وإيصالات السداد لا تُفتح إلا بحساب مصرّح.'),
    ('حذف الحساب', 'الطالب يحذف حسابه وبياناته الشخصية بنفسه من داخل التطبيق.')
]
for i, (t, b) in enumerate(sec):
    col = i % 3
    row = i // 3
    card(s12, M + col * 4.05, 1.85 + row * 2.3, 3.8, 2.05, t, b, icon='🔒', accent=NAVY)

# ══════════════ 13. التطبيق ══════════════
s13 = prs.slides.add_slide(blank_layout)
head(s13, 'تطبيق الهاتف الذكي', 'الوصول')
card(s13, M, 1.85, 5.8, 2.3, 'أندرويد — جاهز تماماً',
     'التطبيق مبني ويعمل، بإشعارات فورية ودخول بالبصمة وحفظ المحاضرات للمذاكرة بلا إنترنت.', icon='▲', accent=GREEN)
card(s13, M + 6.15, 1.85, 5.85, 2.3, 'آيفون — مهيّأ ومبني عبر GitHub Actions',
     'المشروع جاهز ومبني، ومستوفٍ لشروط App Store. يتبقّى حساب المطوّر للنشر الرسمي.', icon='◆', accent=NAVY)

caps = [
    ('العمل بلا إنترنت (Offline)', 'المحاضرات المحفوظة تُقرأ دون اتصال، والصفحات تعمل من الذاكرة المؤقتة.'),
    ('إشعارات النظام اللحظية', 'تصل كإشعار هاتف حقيقي حتى والتطبيق مغلق.'),
    ('تصميم يتكيّف مع كل الشاشات', 'واجهة عربية أصيلة من اليمين لليسار، ووضع ليلي، وتناسق تام.')
]
for i, (t, b) in enumerate(caps):
    card(s13, M + i * 4.05, 4.45, 3.8, 2.1, t, b, icon='◆', accent=BLUE)

# ══════════════ 14. الخلاصة والطلب ══════════════
s14 = prs.slides.add_slide(blank_layout)
add_bg(s14, NAVY)

tx = s14.shapes.add_textbox(Inches(M), Inches(0.6), Inches(11.9), Inches(1.3))
tf = tx.text_frame
tf.word_wrap = True
set_para(tf.paragraphs[0], 'الخلاصة', 13, hex_to_rgb('7FB3E8'), bold=True)
p1 = tf.add_paragraph()
set_para(p1, 'منصة عاملة — والقرارات المطلوبة', 34, WHITE, bold=True)

asks = [
    ('اعتماد التشغيل الرسمي', 'إقرار النظام مرجعاً معتمداً للحضور والواجبات والدرجات بدل الطرق الحالية.'),
    ('تدريب الأساتذة', 'جلسة تعريفية قصيرة لأساتذة المواد على رفع المحاضرات ورصد الحضور والتقييم.'),
    ('حساب مطوّر آبل', 'اشتراك سنوي ٩٩ دولاراً لإتمام نشر تطبيق الآيفون على متجر App Store.'),
    ('مسؤول متابعة البلاغات', 'تكليف موظف بمراجعة بلاغات المحتوى يومياً — شرط بقاء التطبيق في المتاجر.')
]
for i, (t, b) in enumerate(asks):
    col = i % 2
    row = i // 2
    x = M + col * 6.15
    y = 2.05 + row * 2.0
    
    card(s14, x, y, 5.8, 1.8, t, b, icon=str(i + 1),
         tone=hex_to_rgb('164B7D'), border_color=hex_to_rgb('256296'), accent=AMBER, title_color=WHITE)

tx_foot14 = s14.shapes.add_textbox(Inches(M), Inches(6.55), Inches(11.9), Inches(0.4))
set_para(tx_foot14.text_frame.paragraphs[0], 'كلية النصر التقنية — النظام الجامعي الرقمي المتكامل · الإصدار ٢.١.٠', 12, hex_to_rgb('7FA5CC'))

# حفظ الملفات
out_file = 'عرض-نظام-كلية-النصر-التقنية.pptx'
prs.save(out_file)

# نسخ للمجلدات المعتادة
desktop = os.path.expanduser('~/Desktop')
downloads = os.path.expanduser('~/Downloads')

for dest_dir in [desktop, downloads]:
    try:
        if os.path.exists(dest_dir):
            shutil.copy2(out_file, os.path.join(dest_dir, 'عرض-نظام-كلية-النصر-التقنية.pptx'))
            print(f"Copied to: {dest_dir}")
    except Exception as e:
        print(f"Copy error: {e}")

print('DONE! Pure PowerPoint-compliant presentation generated successfully!')
